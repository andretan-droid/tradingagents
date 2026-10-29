#!/usr/bin/env python3
"""Build Alphabet (GOOGL) Q3 2026 earnings update: charts (matplotlib) + DOCX (python-docx).

All reported figures come from Alphabet's 8-K Exhibit 99.1 earnings releases on SEC EDGAR
(Q4'25, Q1'26, Q2'26, Q3'26; prior-year comparatives taken from the same exhibits).
Anything labelled "E" or "our estimate" is an analyst estimate, not company data.
"""
import os

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import FuncFormatter
from docx import Document
from docx.enum.dml import MSO_THEME_COLOR_INDEX
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
CHART_DIR = os.path.join(HERE, "charts")
OUT = os.path.join(HERE, "Alphabet_Q3_2026_Earnings_Update.docx")
os.makedirs(CHART_DIR, exist_ok=True)

# ----------------------------------------------------------------------------
# Sources (all verified to exist on 2026-10-29)
# ----------------------------------------------------------------------------
URL = {
    "q3_pr": "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000080/googexhibit991q32026.htm",
    "q3_8k": "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000080/0001652044-26-000080-index.htm",
    "q3_10q": "https://www.sec.gov/ix?doc=/Archives/edgar/data/1652044/000165204426000084/goog-20260930.htm",
    "q2_pr": "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000066/googexhibit991q22026.htm",
    "q1_pr": "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000043/googexhibit991q12026.htm",
    "q4_pr": "https://www.sec.gov/Archives/edgar/data/1652044/000165204426000012/googexhibit991q42025.htm",
    "call": "https://www.fool.com/earnings/call-transcripts/2026/10/29/alphabet-googl-q3-2026-earnings-call-transcript/",
    "call_inv": "https://www.investing.com/news/transcripts/earnings-call-transcript-alphabet-q3-2026-profit-miss-sparks-10-afterhours-drop-93CH-4981037",
    "slides_inv": "https://www.investing.com/news/company-news/alphabet-q3-2026-slides-128b-revenue-overshadowed-by-profit-miss-93CH-4981067",
    "mb": "https://www.marketbeat.com/earnings/reports/2026-10-28-alphabet-inc-stock/",
    "cnbc": "https://www.cnbc.com/2026/10/28/alphabet-googl-earnings-q3-2026.html",
    "sa_hist": "https://stockanalysis.com/stocks/googl/history/",
    "sa_quote": "https://stockanalysis.com/stocks/googl/",
    "sa_fc": "https://stockanalysis.com/stocks/googl/forecast/",
    "q1_capex": "https://www.spglobal.com/market-intelligence/en/news-insights/research/2026/05/alphabet-postq-snapshot-ai-momentum-drives-cloud-surge-capex-outlook-rises1",
    "q2_capex": "https://www.fierce-network.com/cloud/google-cloud-surges-alphabets-ai-capex-problem-grows",
    "q4_backlog": "https://www.marketbeat.com/originals/alphabet-crushes-earnings-again-reports-48-cloud-revenue-growth/",
    "edgar": "https://www.sec.gov/cgi-bin/browse-edgar?action=getcompany&CIK=1652044&type=10-Q&dateb=&owner=include&count=10",
}

# ----------------------------------------------------------------------------
# Reported quarterly data ($M except EPS) -- 8-K Exhibit 99.1 releases
# ----------------------------------------------------------------------------
Q = ["Q4'24", "Q1'25", "Q2'25", "Q3'25", "Q4'25", "Q1'26", "Q2'26", "Q3'26"]
D = {
    "rev":    [96469, 90234, 96428, 102346, 113828, 109896, 119796, 127577],
    "search": [54034, 50702, 54190, 56567, 63073, 60399, 63271, 63350],
    "yt":     [10473, 8927, 9796, 10261, 11383, 9883, 11055, 11207],
    "net":    [7954, 7256, 7354, 7354, 7828, 6971, 7303, 7455],
    "subs":   [11633, 10379, 11203, 12870, 13578, 12384, 12911, 14290],
    "svc":    [84094, 77264, 82543, 87052, 95862, 89637, 94540, 96302],
    "cloud":  [11955, 12260, 13624, 15157, 17664, 20028, 24768, 30709],
    "oi":     [30972, 30606, 31271, 31228, 35934, 39696, 40770, 40576],
    "svc_oi": [32836, 32682, 33063, 33527, 40132, 40589, 39544, 39157],
    "cl_oi":  [2093, 2177, 2826, 3594, 5313, 6598, 8814, 11121],
    "ni_c":   [26536, 34540, 28196, 34979, 34455, 62578, 112107, 26170],  # NI to common
    "eq":     [364, 9758, 1286, 10734, 2302, 36915, 99031, -8664],  # gain/(loss) equity securities
    "shares": [12348, 12291, 12198, 12203, 12228, 12238, 12309, 12368],
    "eps":    [2.15, 2.81, 2.31, 2.87, 2.82, 5.11, 9.11, 2.12],
    "capex":  [14276, 17197, 22446, 23953, 27851, 35674, 44924, 51355],
    "ocf":    [39113, 36150, 27747, 48414, 52402, 45790, 39069, 42290],
}
TAX = 0.21  # assumed tax rate applied to equity-securities marks (our assumption)
D["eps_x"] = [(n - e * (1 - TAX)) / s for n, e, s in zip(D["ni_c"], D["eq"], D["shares"])]
D["fcf"] = [o - c for o, c in zip(D["ocf"], D["capex"])]

# Cloud backlog ($B) -- Q3'26 from call transcript; earlier quarters from prior calls via press coverage
BACKLOG_Q = ["Q3'25", "Q4'25", "Q1'26", "Q2'26", "Q3'26"]
BACKLOG = [155, 240, 462, 514, 561]

# Estimates ($B except EPS) -- OUR estimates
EST = {
    #               prior FY26E, new FY26E, prior FY27E, new FY27E
    "Revenue":         (496.6, 495.7, 612.6, 616.2),
    "Google Services": (388.8, 384.5, 431.0, 422.9),
    "Google Cloud":    (106.3, 109.5, 180.0, 191.6),
    "Operating income": (167.0, 163.9, 199.5, 194.3),
    "EPS ex-MTM ($)":  (11.58, 10.98, 13.12, 12.72),
}

PRICE_CLOSE = 324.89
PRICE_PRE = 291.08
PT = 360

# ----------------------------------------------------------------------------
# Chart styling
# ----------------------------------------------------------------------------
C1, C2, C3, C4 = "#2a78d6", "#eb6834", "#1baf7a", "#eda100"
GRAY, INK, INK2 = "#b9b8b3", "#0b0b0b", "#52514e"
plt.rcParams.update({
    "font.family": "serif",
    "font.serif": ["Times New Roman", "Liberation Serif", "DejaVu Serif"],
    "font.size": 9.5,
    "axes.edgecolor": GRAY,
    "axes.labelcolor": INK2,
    "xtick.color": INK2,
    "ytick.color": INK2,
    "axes.spines.top": False,
    "axes.spines.right": False,
    "axes.grid": True,
    "axes.grid.axis": "y",
    "grid.color": "#e6e5e0",
    "grid.linewidth": 0.6,
    "legend.frameon": False,
    "figure.dpi": 200,
})
W, H = 6.4, 2.75


def save(fig, name):
    path = os.path.join(CHART_DIR, name)
    fig.tight_layout()
    fig.savefig(path, dpi=200, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


def bfmt(x, _):
    return f"${x:,.0f}B"


def chart_segments():
    fig, ax = plt.subplots(figsize=(W, H))
    x = range(len(Q))
    svc = [v / 1000 for v in D["svc"]]
    cl = [v / 1000 for v in D["cloud"]]
    oth = [(r - s - c) / 1000 for r, s, c in zip(D["rev"], D["svc"], D["cloud"])]
    ax.bar(x, svc, color=C1, width=0.62, label="Google Services", edgecolor="white", linewidth=1)
    ax.bar(x, cl, bottom=svc, color=C2, width=0.62, label="Google Cloud", edgecolor="white", linewidth=1)
    bot = [a + b for a, b in zip(svc, cl)]
    ax.bar(x, oth, bottom=bot, color=GRAY, width=0.62, label="Other Bets + hedging", edgecolor="white", linewidth=1)
    for i, r in enumerate(D["rev"]):
        ax.text(i, r / 1000 + 2, f"${r/1000:.1f}B", ha="center", va="bottom", fontsize=8.5, color=INK)
    ax.set_xticks(list(x))
    ax.set_xticklabels(Q)
    ax.yaxis.set_major_formatter(FuncFormatter(bfmt))
    ax.set_ylim(0, 145)
    ax.legend(loc="upper left", ncol=3, fontsize=8.5)
    return save(fig, "fig01_revenue_by_segment.png")


def chart_cloud():
    fig, (a1, a2) = plt.subplots(1, 2, figsize=(W, H))
    x = range(len(Q))
    cl = [v / 1000 for v in D["cloud"]]
    a1.bar(x, cl, color=C2, width=0.62)
    for i, v in enumerate(cl):
        a1.text(i, v + 0.5, f"{v:.1f}", ha="center", fontsize=7.5, color=INK)
    a1.set_title("Google Cloud revenue ($B)", fontsize=9.5, color=INK)
    a1.set_xticks(list(x))
    a1.set_xticklabels(Q, rotation=45, fontsize=7.5)
    m = [o / r * 100 for o, r in zip(D["cl_oi"], D["cloud"])]
    a2.plot(list(x), m, color=C2, linewidth=2, marker="o", markersize=5)
    for i, v in enumerate(m):
        a2.text(i, v + 1.5, f"{v:.1f}%", ha="center", fontsize=7.5, color=INK)
    a2.set_title("Google Cloud operating margin (%)", fontsize=9.5, color=INK)
    a2.set_ylim(0, 45)
    a2.set_xlim(-0.7, 7.5)
    a2.set_xticks(list(x))
    a2.set_xticklabels(Q, rotation=45, fontsize=7.5)
    return save(fig, "fig02_cloud_revenue_margin.png")


def chart_backlog():
    fig, ax = plt.subplots(figsize=(W, 2.4))
    x = range(len(BACKLOG_Q))
    ax.bar(x, BACKLOG, color=C2, width=0.55)
    for i, v in enumerate(BACKLOG):
        ax.text(i, v + 10, f"${v}B", ha="center", fontsize=8.5, color=INK)
    ax.set_xticks(list(x))
    ax.set_xticklabels(BACKLOG_Q)
    ax.yaxis.set_major_formatter(FuncFormatter(bfmt))
    ax.set_ylim(0, 650)
    return save(fig, "fig03_cloud_backlog.png")


def chart_margins():
    fig, ax = plt.subplots(figsize=(W, H))
    x = list(range(len(Q)))
    series = [
        ("Google Services", [o / r * 100 for o, r in zip(D["svc_oi"], D["svc"])], C1),
        ("Alphabet consolidated", [o / r * 100 for o, r in zip(D["oi"], D["rev"])], INK2),
        ("Google Cloud", [o / r * 100 for o, r in zip(D["cl_oi"], D["cloud"])], C2),
    ]
    for name, s, c in series:
        ax.plot(x, s, color=c, linewidth=2, marker="o", markersize=4.5, label=name)
        ax.text(x[-1] + 0.15, s[-1], f"{s[-1]:.1f}%", va="center", fontsize=8.5, color=INK)
    ax.set_xticks(x)
    ax.set_xticklabels(Q)
    ax.set_xlim(-0.3, len(Q) - 0.3)
    ax.set_ylim(10, 50)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
    ax.legend(loc="lower right", ncol=3, fontsize=8.5)
    return save(fig, "fig04_operating_margins.png")


def chart_eps():
    fig, ax = plt.subplots(figsize=(W, H))
    x = list(range(len(Q)))
    w = 0.38
    ax.bar([i - w / 2 for i in x], D["eps"], width=w, color=GRAY, label="GAAP diluted EPS", edgecolor="white")
    ax.bar([i + w / 2 for i in x], D["eps_x"], width=w, color=C1, label="EPS ex-equity securities marks (our calc.)", edgecolor="white")
    for i in x:
        ax.text(i - w / 2, D["eps"][i] - 0.45, f"{D['eps'][i]:.2f}", ha="center", fontsize=7, color=INK)
        ax.text(i + w / 2, D["eps_x"][i] - 0.45, f"{D['eps_x'][i]:.2f}", ha="center", fontsize=7, color="white")
    ax.axhline(3.01, color=C2, linewidth=1.2, linestyle="--", label="Q3'26 consensus EPS ($3.01)")
    ax.set_xticks(x)
    ax.set_xticklabels(Q)
    ax.set_ylim(0, 10)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:.0f}"))
    ax.legend(loc="upper left", fontsize=8.5)
    return save(fig, "fig05_eps_gaap_vs_ex_marks.png")


def chart_bridge():
    # Q3'25 GAAP EPS -> Q3'26 GAAP EPS, per-share, after tax at 21% (our analysis)
    s = 12368
    steps = [
        ("Q3'25 EPS", 2.87, "total"),
        ("Operating\nincome", (40576 - 31228) * (1 - TAX) / s, "d"),
        ("Equity securities\nswing", (-8664 - 10734) * (1 - TAX) / s, "d"),
        ("Other OI&E\n(interest, debt sec.)", ((-8221 + 8664) - (12759 - 10734)) * (1 - TAX) / s, "d"),
        ("Preferred\ndividends", -298 / s, "d"),
    ]
    run = sum(v for _, v, _ in steps)
    steps.append(("Tax rate, share\ncount, other", 2.12 - run, "d"))
    steps.append(("Q3'26 EPS", 2.12, "total"))
    fig, ax = plt.subplots(figsize=(W, 2.9))
    cum = 0
    for i, (lab, v, kind) in enumerate(steps):
        if kind == "total":
            ax.bar(i, v, color=INK2, width=0.6)
            cum = v
            ax.text(i, v + 0.05, f"${v:.2f}", ha="center", fontsize=8.5, color=INK)
        else:
            base = cum if v >= 0 else cum + v
            ax.bar(i, abs(v), bottom=base, color=C3 if v >= 0 else C2, width=0.6)
            ax.text(i, max(cum, cum + v) + 0.05, f"{v:+.2f}", ha="center", fontsize=8.5, color=INK)
            cum += v
    ax.set_xticks(range(len(steps)))
    ax.set_xticklabels([s_[0] for s_ in steps], fontsize=7.6)
    ax.set_ylim(0, 3.8)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:.2f}"))
    return save(fig, "fig06_eps_bridge.png")


def chart_cash():
    fig, ax = plt.subplots(figsize=(W, H))
    x = list(range(len(Q)))
    w = 0.27
    ax.bar([i - w for i in x], [v / 1000 for v in D["ocf"]], width=w, color=C1, label="Operating cash flow")
    ax.bar(x, [v / 1000 for v in D["capex"]], width=w, color=C2, label="Capex")
    ax.bar([i + w for i in x], [v / 1000 for v in D["fcf"]], width=w, color=C3, label="Free cash flow")
    for i in x:
        f = D["fcf"][i] / 1000
        ax.text(i + w, f + (1.2 if f >= 0 else -4.2), f"{f:.1f}", ha="center", fontsize=7, color=INK)
    ax.axhline(0, color=INK2, linewidth=0.8)
    ax.set_xticks(x)
    ax.set_xticklabels(Q)
    ax.yaxis.set_major_formatter(FuncFormatter(bfmt))
    ax.set_ylim(-15, 62)
    ax.legend(loc="upper left", ncol=3, fontsize=8.5)
    return save(fig, "fig07_ocf_capex_fcf.png")


def chart_capex_guide():
    fig, ax = plt.subplots(figsize=(W, 2.5))
    labels = ["Feb-26 (Q4'25 call)", "Apr-26 (Q1'26 call)", "Jul-26 (Q2'26 call)", "Oct-26 (Q3'26 call)"]
    lo = [175, 180, 195, 195]
    hi = [185, 190, 205, 205]
    for i, (a, b) in enumerate(zip(lo, hi)):
        ax.barh(i, b - a, left=a, color=C2, height=0.5)
        ax.text(b + 1, i, f"${a}-{b}B" + ("  (maintained)" if i == 3 else ""), va="center", fontsize=8.5, color=INK)
    ax.set_yticks(range(4))
    ax.set_yticklabels(labels, fontsize=8.5)
    ax.invert_yaxis()
    ax.set_xlim(160, 230)
    ax.grid(axis="x", color="#e6e5e0")
    ax.grid(axis="y", visible=False)
    ax.xaxis.set_major_formatter(FuncFormatter(bfmt))
    ax.set_xlabel("2026 capex guidance range (2025 actual capex: $91.4B)", fontsize=8.5)
    return save(fig, "fig08_capex_guidance.png")


def chart_services_growth():
    fig, ax = plt.subplots(figsize=(W, H))
    qq = Q[4:]
    x = list(range(4))
    lines = [("Search & other", "search", C1), ("YouTube ads", "yt", C2),
             ("Subscriptions, platforms & devices", "subs", C3), ("Network", "net", C4)]
    nudge = {"search": 1.0, "subs": -0.9, "yt": -0.6, "net": 0}
    for name, k, c in lines:
        g = [(D[k][i] / D[k][i - 4] - 1) * 100 for i in range(4, 8)]
        ax.plot(x, g, color=c, linewidth=2, marker="o", markersize=4.5, label=name)
        ax.text(3.1, g[-1] + nudge[k], f"{g[-1]:.1f}%", va="center", fontsize=8.5, color=INK)
    ax.axhline(0, color=INK2, linewidth=0.8)
    ax.set_xticks(x)
    ax.set_xticklabels(qq)
    ax.set_xlim(-0.2, 3.5)
    ax.set_ylim(-6, 24)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:.0f}%"))
    ax.legend(loc="upper center", bbox_to_anchor=(0.5, -0.12), ncol=4, fontsize=8.2)
    return save(fig, "fig09_services_yoy_growth.png")


def chart_revisions():
    fig, axes = plt.subplots(1, 3, figsize=(W, 2.6))
    items = [("Revenue ($B)", "Revenue", "{:.0f}"), ("Operating income ($B)", "Operating income", "{:.0f}"),
             ("EPS ex-MTM ($)", "EPS ex-MTM ($)", "{:.2f}")]
    for ax, (title, k, fmt) in zip(axes, items):
        p26, n26, p27, n27 = EST[k]
        xs = [0, 1]
        ax.bar([i - 0.18 for i in xs], [p26, p27], width=0.34, color=GRAY, label="Prior")
        ax.bar([i + 0.18 for i in xs], [n26, n27], width=0.34, color=C1, label="New")
        for i, (p, n) in enumerate([(p26, n26), (p27, n27)]):
            ax.text(i - 0.18, p * 1.01, fmt.format(p), ha="center", fontsize=6.8, color=INK)
            ax.text(i + 0.18, n * 1.01, fmt.format(n), ha="center", fontsize=6.8, color=INK)
        ax.set_xticks(xs)
        ax.set_xticklabels(["FY26E", "FY27E"])
        ax.set_title(title, fontsize=9, color=INK)
        ax.set_ylim(0, max(p27, n27) * 1.18)
        ax.tick_params(axis="y", labelsize=7.5)
    axes[0].legend(loc="upper left", fontsize=7.5)
    return save(fig, "fig10_estimate_revisions.png")


def chart_sotp():
    parts = [("Google Services\n(22x NOPAT)", 3109), ("Google Cloud\n(35x NOPAT)", 1694),
             ("Other Bets\n(nominal)", 50), ("Alphabet-level\ncosts (22x)", -538),
             ("Net cash\n(9/30/26)", 140)]
    sh = 12.45
    fig, ax = plt.subplots(figsize=(W, 2.8))
    cum = 0
    for i, (lab, v) in enumerate(parts):
        ps = v / sh
        base = cum if ps >= 0 else cum + ps
        ax.bar(i, abs(ps), bottom=base, color=C3 if ps >= 0 else C2, width=0.6)
        ax.text(i, max(cum, cum + ps) + 6, f"{ps:+.0f}", ha="center", fontsize=8.5, color=INK)
        cum += ps
    ax.bar(len(parts), cum, color=C1, width=0.6)
    ax.text(len(parts), cum + 6, f"${cum:.0f}", ha="center", fontsize=8.5, color=INK)
    ax.axhline(PRICE_PRE, color=INK2, linestyle="--", linewidth=1)
    ax.text(-0.45, PRICE_PRE + 6, f"Pre-market price ${PRICE_PRE:.2f}", fontsize=7.5, color=INK2, ha="left")
    ax.set_xticks(range(len(parts) + 1))
    ax.set_xticklabels([p[0] for p in parts] + ["SOTP value\nper share"], fontsize=7.6)
    ax.set_ylim(0, 430)
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"${v:.0f}"))
    return save(fig, "fig11_sotp.png")


# ----------------------------------------------------------------------------
# DOCX helpers
# ----------------------------------------------------------------------------
FONT = "Times New Roman"
BLUE = RGBColor(0x05, 0x63, 0xC1)
NAVY = RGBColor(0x1F, 0x3A, 0x5F)


def set_run_font(run, size=None, bold=None, italic=None, color=None):
    run.font.name = FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def add_hyperlink(paragraph, text, url, size=None, bold=False, italic=False):
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement("w:hyperlink")
    h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(a), FONT)
    rpr.append(rfonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    c = OxmlElement("w:color")
    c.set(qn("w:val"), "0563C1")
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    if size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(int(size * 2)))
        rpr.append(sz)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    h.append(r)
    paragraph._p.append(h)
    return h


def para(doc, text="", size=10.5, bold=False, italic=False, align=None, space_after=4, space_before=0,
         color=None, keep_next=False):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    pf.line_spacing = 1.0
    if keep_next:
        pf.keep_with_next = True
    if align:
        p.alignment = align
    if text:
        set_run_font(p.add_run(text), size, bold, italic, color)
    return p


def rich(doc, segments, size=10.5, space_after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """segments: list of str | ('b', str) | ('i', str) | ('link', text, url)."""
    p = para(doc, size=size, space_after=space_after, align=align)
    for s in segments:
        if isinstance(s, str):
            set_run_font(p.add_run(s), size)
        elif s[0] == "b":
            set_run_font(p.add_run(s[1]), size, bold=True)
        elif s[0] == "i":
            set_run_font(p.add_run(s[1]), size, italic=True)
        elif s[0] == "link":
            add_hyperlink(p, s[1], s[2], size=size)
    return p


def heading(doc, text, size=13, space_before=8):
    p = para(doc, text, size=size, bold=True, color=NAVY, space_before=space_before, space_after=4, keep_next=True)
    # bottom border
    pbdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), "6")
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), "1F3A5F")
    pbdr.append(b)
    p._p.get_or_add_pPr().append(pbdr)
    return p


def subheading(doc, text):
    return para(doc, text, size=11, bold=True, color=NAVY, space_before=5, space_after=3, keep_next=True)


def bullet(doc, header, body_segments, size=10.5):
    p = para(doc, size=size, space_after=5, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
    set_run_font(p.add_run("■ "), size, color=NAVY)
    set_run_font(p.add_run(header + " "), size, bold=True)
    for s in body_segments:
        if isinstance(s, str):
            set_run_font(p.add_run(s), size)
        elif s[0] == "link":
            add_hyperlink(p, s[1], s[2], size=size)
        elif s[0] == "b":
            set_run_font(p.add_run(s[1]), size, bold=True)
    return p


def source_line(doc, segments, size=8.5):
    p = para(doc, size=size, space_after=7, italic=True)
    set_run_font(p.add_run("Source: "), size, italic=True)
    for s in segments:
        if isinstance(s, str):
            set_run_font(p.add_run(s), size, italic=True)
        else:
            add_hyperlink(p, s[1], s[2], size=size, italic=True)
    return p


FIG = [0]


def figure(doc, title, path, src, width=6.3):
    FIG[0] += 1
    # caption and picture share one paragraph so they can never be split across pages
    p = para(doc, space_before=4, space_after=1, keep_next=True)
    set_run_font(p.add_run(f"Figure {FIG[0]} - {title}"), 9.5, bold=True)
    r = p.add_run()
    r.add_break()
    r.add_picture(path, width=Inches(width))
    source_line(doc, src)


TBL = [0]


def shade(cell, hexcolor):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hexcolor)
    tcpr.append(shd)


def table(doc, title, header, rows, src, col_widths=None, size=8.8, bold_rows=()):
    TBL[0] += 1
    para(doc, f"Table {TBL[0]} - {title}", size=9.5, bold=True, space_before=4, space_after=2, keep_next=True)
    t = doc.add_table(rows=1 + len(rows), cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]
        c.text = ""
        r = c.paragraphs[0].add_run(h)
        set_run_font(r, size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER if j else WD_ALIGN_PARAGRAPH.LEFT
        c.paragraphs[0].paragraph_format.keep_with_next = True
        shade(c, "1F3A5F")
    for i, row in enumerate(rows, start=1):
        for j, v in enumerate(row):
            c = t.rows[i].cells[j]
            c.text = ""
            r = c.paragraphs[0].add_run(str(v))
            set_run_font(r, size, bold=(i - 1) in bold_rows)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT if j else WD_ALIGN_PARAGRAPH.LEFT
            c.paragraphs[0].paragraph_format.space_after = Pt(0)
            if i < len(rows):
                c.paragraphs[0].paragraph_format.keep_with_next = True
            if i % 2 == 0:
                shade(c, "EEF2F7")
    if col_widths:
        t.autofit = False
        for j, w in enumerate(col_widths):
            t.columns[j].width = Inches(w)
        for row in t.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Inches(w)
    source_line(doc, src)
    return t


# ----------------------------------------------------------------------------
# Build
# ----------------------------------------------------------------------------
def build():
    figs = {
        "seg": chart_segments(), "cloud": chart_cloud(), "backlog": chart_backlog(), "margins": chart_margins(),
        "eps": chart_eps(), "bridge": chart_bridge(), "cash": chart_cash(), "guide": chart_capex_guide(),
        "svcg": chart_services_growth(), "rev": chart_revisions(), "sotp": chart_sotp(),
    }

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(sec, side, Inches(0.9))
    sec.top_margin = Inches(0.75)
    sec.bottom_margin = Inches(0.75)
    st = doc.styles["Normal"]
    st.font.name = FONT
    st.font.size = Pt(10.5)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    # Footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(fp.add_run("Alphabet Inc. (GOOGL) | Q3 2026 Earnings Update | October 29, 2026 | "
                            "Estimates are analyst estimates, not company guidance."), 8, italic=True, color=RGBColor(0x52, 0x51, 0x4E))

    # ---------------- PAGE 1 ----------------
    para(doc, "ALPHABET INC. (NASDAQ: GOOGL)", size=17, bold=True, color=NAVY, space_after=0)
    para(doc, "Q3 2026 EARNINGS UPDATE  |  October 29, 2026  |  Internet / Cloud Infrastructure", size=10.5, bold=True,
         color=RGBColor(0x52, 0x51, 0x4E), space_after=4)
    para(doc, "Cloud Doubles, SpaceX Mark Sinks EPS; Capex Is the Debate - Maintain BUY, PT $360",
         size=14.5, bold=True, space_after=4)
    rich(doc, [("b", "Rating: MAINTAIN BUY   |   12-month price target: $360"),
               " (set in this note; 50% P/E, 50% sum-of-the-parts)   |   ",
               ("b", f"Price: ${PRICE_PRE:.2f} pre-market 10/29 "), f"(close 10/28: ${PRICE_CLOSE:.2f})   |   ",
               ("b", f"Implied upside: +{(PT / PRICE_PRE - 1) * 100:.0f}%"),
               "   |   52-week range: $266.50-$408.61   |   Market cap: ~$3.6T at pre-market price"],
         size=9.5, space_after=6, align=WD_ALIGN_PARAGRAPH.LEFT)

    table(doc, "Q3 2026 Earnings Summary: Revenue In Line, EPS Miss (Non-Operating)",
          ["$M except EPS", "Q3'26A", "Consensus", "Variance", "Q3'25A", "YoY"],
          [
              ["Revenue", "127,577", "127,970", "-393 (-0.3%)", "102,346", "+24.7%"],
              ["Google Cloud revenue", "30,709", "n/a", "n/a", "15,157", "+102.6%"],
              ["Operating income", "40,576", "n/a", "n/a", "31,228", "+29.9%"],
              ["Operating margin", "31.8%", "n/a", "n/a", "30.5%", "+130bps"],
              ["Diluted EPS (GAAP)", "$2.12", "$3.01", "-$0.89 (-29.6%)", "$2.87", "-26.1%"],
              ["EPS ex-equity securities marks*", "$2.67", "$3.01", "-$0.34 (-11.3%)", "$2.17", "+23.0%"],
              ["Capex", "51,355", "n/a", "n/a", "23,953", "+114.4%"],
              ["Free cash flow", "(9,065)", "n/a", "n/a", "24,461**", "n.m."],
          ],
          [("link", "Q3'26 earnings release (8-K Ex. 99.1, Oct 28, 2026)", URL["q3_pr"]),
           "; consensus from ", ("link", "MarketBeat (pre-print, accessed Oct 29, 2026)", URL["mb"]),
           "; other outlets cited $3.02 EPS and $127.2-128.2B revenue. *Our calculation: removes the $8,664M equity "
           "securities loss (and prior-year gains) at an assumed 21% tax rate. **Q3'25 FCF derived as OCF less capex."],
          col_widths=[2.0, 0.85, 0.85, 1.15, 0.85, 0.8], bold_rows=(0, 4))

    para(doc, "Key takeaways", size=11, bold=True, color=NAVY, space_after=2, keep_next=True)
    bullet(doc, "Operations beat; the headline miss is a SpaceX mark.",
           ["Operating income rose 30% to $40.6B (31.8% margin), but other income swung to a $8.2B loss, driven mainly "
            "by an unrealized loss on the SpaceX stake (an $8,664M equity securities loss). That item alone cost about "
            "$0.55 of EPS. Excluding marks, EPS was ~$2.67, up 23% YoY but still ~11% below the $3.01 consensus "
            "because of higher interest expense, preferred dividends and a larger share count."])
    bullet(doc, "Google Cloud is now a second profit engine.",
           ["Cloud revenue doubled (+103%) to $30.7B, with operating income of $11.1B (36.2% margin vs. 23.7%). "
            "Backlog grew to $561B, and management expects just over 55% of it to become revenue within 24 months."])
    bullet(doc, "Capex and negative FCF are the debate.",
           ["Capex of $51.4B (+114%) pushed FCF to -$9.1B. 2026 guidance stays at $195-205B, and 2027 capex is "
            "expected to \"increase significantly.\" The shares fell ~10% pre-market."])

    # ---------------- INVESTMENT IMPACT ----------------
    heading(doc, "Investment Impact")
    bullet(doc, "Revenue was in line; mix moved further toward Cloud.",
           ["Revenue of $127.6B (+25% YoY, +25% constant currency) was $0.4B (0.3%) below the MarketBeat consensus "
            "of $128.0B and about 0.3% above the $127.2B figure cited by Investing.com. Either way, it was in line. "
            "Mix was the story. Cloud added $15.6B of YoY revenue, more than Google Services (+$9.3B, +11%). Search & "
            "other grew 12% to $63.4B, slowing from +17% in Q2. Management attributed most of the slowdown to FX, "
            "EU de minimis rule changes hitting APAC retail advertisers, and lapping World Cup media spend. Cloud is "
            "now 24% of revenue vs. 15% a year ago. Its incremental margin was 48% ($7.5B of added operating income "
            "on $15.6B of added revenue), which drove the +130bps consolidated margin expansion."])
    bullet(doc, "The EPS miss is mostly non-operating, but not entirely.",
           ["GAAP diluted EPS of $2.12 missed the $3.01 consensus by $0.89. The $8,664M equity securities loss explains "
            "about $0.55. The other ~$0.34 came from items we expect to persist. Interest expense rose to $967M "
            "(vs. $143M) after long-term debt grew to $125.1B. Preferred dividends of $298M began on the June mandatory "
            "convertible issue. Diluted shares rose 1.4% YoY to 12,368M after at-the-market (ATM) issuance. The "
            "\"Alphabet-level activities\" loss also widened to $7.7B from $4.5B. In our view, the market is right to "
            "treat the mark as noise and wrong to ignore the financing costs. Both are now in our model."])
    bullet(doc, "Capex guidance was held, but 2027 is a step up and FCF turned negative again.",
           ["Q3 capex of $51.4B took year-to-date spend to $132.0B. That leaves $63-73B for Q4 within the unchanged "
            "$195-205B range. CFO Anat Ashkenazi repeated that \"we continue to expect our CapEx to increase "
            "significantly in 2027.\" Operating cash flow fell 13% YoY to $42.3B. Management cited inventory built for "
            "TPU system sales and a $5.2B cash payment of the 2018 European Commission fine. Trailing-12-month FCF is "
            "down to $19.7B. Alphabet has bought back no stock in 2026, and the balance sheet is funding the build."])
    bullet(doc, "Maintain BUY with a $360 PT; the sell-off over-discounts a non-cash mark.",
           [f"At ${PRICE_PRE:.0f}, GOOGL trades at 22.9x our 2027E EPS ex-marks of $12.72, or 26.5x 2026E. That is a "
            "modest multiple for a company growing revenue 25% with a Cloud business that has doubled and a $561B "
            "backlog. Our $360 target blends 28x 2027E EPS ($356) with a sum-of-the-parts ($358) and implies +24% "
            "upside. The main risk is that capex keeps rising faster than Cloud revenue, leaving FCF negative into "
            "2027. That risk is real, but the Q3 Cloud margin of 36% shows the spend is earning a return today."])

    table(doc, "Updated Financial Estimates: Prior vs. New (our estimates)",
          ["$B except EPS", "FY26E Prior", "FY26E New", "Chg", "FY27E Prior", "FY27E New", "Chg"],
          [
              ["Revenue", "496.6", "495.7", "-0.2%", "612.6", "616.2", "+0.6%"],
              ["  Google Services", "388.8", "384.5", "-1.1%", "431.0", "422.9", "-1.9%"],
              ["  Google Cloud", "106.3", "109.5", "+3.0%", "180.0", "191.6", "+6.4%"],
              ["Operating income", "167.0", "163.9", "-1.9%", "199.5", "194.3", "-2.6%"],
              ["Operating margin", "33.6%", "33.1%", "-50bps", "32.6%", "31.5%", "-110bps"],
              ["EPS ex-equity marks ($)", "11.58", "10.98", "-5.2%", "13.12", "12.72", "-3.0%"],
              ["GAAP EPS ($)*", "20.32", "19.16", "-5.7%", "13.12", "12.72", "-3.0%"],
              ["Capex", "200.0", "200.0", "0.0%", "250.0", "270.0", "+8.0%"],
              ["Free cash flow", "(15.0)", "(22.9)", "n.m.", "(20.0)", "(40.0)", "n.m."],
              [f"P/E on EPS ex-marks @ ${PRICE_PRE:.0f}", "25.1x", "26.5x", "", "22.2x", "22.9x", ""],
          ],
          ["Our estimates. \"Prior\" is our pre-print model, which was aligned with Street Q3 consensus ($128.0B revenue, "
           "$3.01 EPS). FY26 includes 9M'26 actuals from the ",
           ("link", "Q3'26 release", URL["q3_pr"]),
           ". *GAAP FY26E includes $127.3B of year-to-date equity securities gains. We forecast no future marks. "
           "Street FY26/FY27 EPS of $11.89/$15.19 (", ("link", "StockAnalysis, Oct 29, 2026", URL["sa_fc"]),
           ") may not be on a comparable basis."],
          col_widths=[1.85, 0.8, 0.8, 0.7, 0.8, 0.8, 0.75], bold_rows=(0, 5))

    # ---------------- RESULTS ANALYSIS ----------------
    heading(doc, "Detailed Results Analysis")
    subheading(doc, "Revenue: Cloud doubles while Services slows on FX and tough comps")
    rich(doc, ["Total revenue grew 24.7% YoY to $127,577M, Alphabet's 13th straight quarter of double-digit growth. "
               "Sequential growth was 6.5% (7% constant currency). Google Services revenue grew 10.6% to $96,302M, "
               "the slowest of the last four quarters. Search & other rose 12.0% to $63,350M (Q2: +16.8%; "
               "Q1: +19.1%). YouTube ads rose 9.2% to $11,207M (Q2: +12.9%). Subscriptions, platforms & devices "
               "rose 11.0% to $14,290M, and Network rose 1.4% to $7,455M. Management said the Search slowdown came "
               "mostly from FX, EU de minimis changes and lapping World Cup spend. It said Search query growth was "
               "the highest in five years and that AI Mode queries are on average 3x longer than traditional "
               "searches. Q4 has a tougher comparison (Q4'25 Search +16.7%), and the CFO guided to \"a slightly "
               "larger\" FX headwind. We model Q4 Search at +9%."])
    figure(doc, "Quarterly revenue by segment ($B), Q4'24-Q3'26", figs["seg"],
           [("link", "Alphabet 8-K Ex. 99.1 releases", URL["q3_pr"]), " for ",
            ("link", "Q4'25", URL["q4_pr"]), ", ", ("link", "Q1'26", URL["q1_pr"]), ", ",
            ("link", "Q2'26", URL["q2_pr"]), " and ", ("link", "Q3'26", URL["q3_pr"]),
            " (prior-year quarters from the same releases)."])

    table(doc, "Quarterly Segment Progression ($M)",
          ["$M", "Q3'25", "Q4'25", "Q1'26", "Q2'26", "Q3'26", "YoY", "QoQ"],
          [
              ["Search & other", "56,567", "63,073", "60,399", "63,271", "63,350", "+12.0%", "+0.1%"],
              ["YouTube ads", "10,261", "11,383", "9,883", "11,055", "11,207", "+9.2%", "+1.4%"],
              ["Google Network", "7,354", "7,828", "6,971", "7,303", "7,455", "+1.4%", "+2.1%"],
              ["Subs, platforms & devices", "12,870", "13,578", "12,384", "12,911", "14,290", "+11.0%", "+10.7%"],
              ["Google Services", "87,052", "95,862", "89,637", "94,540", "96,302", "+10.6%", "+1.9%"],
              ["Google Cloud", "15,157", "17,664", "20,028", "24,768", "30,709", "+102.6%", "+24.0%"],
              ["Total revenue", "102,346", "113,828", "109,896", "119,796", "127,577", "+24.7%", "+6.5%"],
              ["Services op. income", "33,527", "40,132", "40,589", "39,544", "39,157", "+16.8%", "-1.0%"],
              ["Cloud op. income", "3,594", "5,313", "6,598", "8,814", "11,121", "+209.4%", "+26.2%"],
              ["Total op. income", "31,228", "35,934", "39,696", "40,770", "40,576", "+29.9%", "-0.5%"],
          ],
          [("link", "Alphabet Q3'26", URL["q3_pr"]), ", ", ("link", "Q2'26", URL["q2_pr"]), ", ",
           ("link", "Q1'26", URL["q1_pr"]), " and ", ("link", "Q4'25", URL["q4_pr"]),
           " earnings releases (8-K Ex. 99.1). Q4'25 operating income includes a $2.1B Waymo compensation charge."],
          col_widths=[1.75, 0.68, 0.68, 0.68, 0.68, 0.68, 0.7, 0.65], bold_rows=(4, 5, 6, 9))

    figure(doc, "Google Services revenue growth YoY by line, Q4'25-Q3'26", figs["svcg"],
           [("link", "Alphabet 8-K Ex. 99.1 releases", URL["q3_pr"]), " (Q4'24-Q3'26); our calculations."])

    subheading(doc, "Profitability: margin up 130bps despite a 49% jump in depreciation")
    rich(doc, ["Operating income rose 29.9% to $40,576M, and the margin widened 130bps to 31.8%. That came despite "
               "depreciation of $8,369M (+49% YoY), R&D up 36% to $20,640M, and headcount up 7.7% to 204,757. "
               "G&A fell 17% to $6,138M because Q3'25 carried the European Commission fine. Google Services margin "
               "was 40.7% (Q3'25: 38.5%), down from 41.8% in Q2, partly reflecting rising "
               "depreciation. Cloud margin rose to 36.2% from 35.6% in Q2 and 23.7% a year ago. Other Bets lost "
               "$2,047M (vs. $1,426M), and Alphabet-level activities, which include some AI model development costs, "
               "cost $7,655M (vs. $4,467M). Ashkenazi flagged near-term Cloud margin \"headwinds\" as third-party "
               "compute capacity, which came online at the end of Q3, ramps in Q4. We model Cloud margin at 32% in "
               "Q4 and 31% in 2027."])
    figure(doc, "Operating margin by segment, Q4'24-Q3'26", figs["margins"],
           [("link", "Alphabet 8-K Ex. 99.1 releases", URL["q3_pr"]), "; our calculations."])

    subheading(doc, "Below the line: the SpaceX mark and rising financing costs")
    rich(doc, ["Other income (expense) was -$8,221M vs. +$12,759M in Q3'25. Components: a $8,664M loss on equity "
               "securities (Q3'25: +$10,734M gain), a $383M loss on debt securities, interest income of $1,768M and "
               "interest expense of $967M. Management attributed the loss \"primarily\" to the unrealized decline in "
               "the SpaceX investment. Marks have flattered GAAP EPS all year: equity gains were $36.9B in Q1 and "
               "$99.0B in Q2, lifting Q2 EPS to $9.11. Alphabet now holds $90.4B of marketable equity securities, "
               "so these swings will continue. The tax rate was 18.2%. Net income available to common shareholders "
               "fell 25% to $26,170M after the first $298M of preferred dividends. On our ex-marks measure, EPS rose "
               "23% YoY to $2.67, in line with Q4'25-Q2'26 ($2.67-$2.75). The underlying business is steady, but it "
               "did not grow EPS sequentially because depreciation and financing costs absorbed the gains from Cloud."])
    figure(doc, "Diluted EPS: GAAP vs. ex-equity securities marks, Q4'24-Q3'26", figs["eps"],
           [("link", "Alphabet 8-K Ex. 99.1 releases", URL["q3_pr"]),
            "; ex-marks EPS is our calculation (21% tax on marks); consensus from ",
            ("link", "MarketBeat", URL["mb"]), "."])
    figure(doc, "Q3 EPS bridge, Q3'25 to Q3'26 ($ per diluted share, after tax)", figs["bridge"],
           [("link", "Q3'26 earnings release", URL["q3_pr"]), "; our analysis at an assumed 21% tax rate on pre-tax items."])

    # ---------------- KEY METRICS & GUIDANCE ----------------
    heading(doc, "Key Metrics & Guidance")
    subheading(doc, "Cloud: backlog of $561B, TPU system sales, and 36% margins")
    rich(doc, ["Google Cloud revenue of $30,709M grew 102.6% YoY and 24.0% QoQ, the fastest growth in our data series. "
               "Ashkenazi said \"Google Cloud's backlog was $561 billion in the third quarter,\" up 9% QoQ from "
               "$514B and 3.6x the ~$155B of Q3'25. Just over 55% is expected to convert to revenue within 24 months. "
               "That is roughly $309B, or about $154B a year, before new bookings. TPU system sales (customers buying "
               "Google's TPU systems rather than renting capacity) were a \"meaningful\" contributor. Management "
               "expects them \"to continue to scale in Q4, followed by a meaningful revenue ramp in 2027.\" Other "
               "disclosures: more than 2,600 customers each used over 100 billion tokens in the past 12 months, "
               "nearly 90% of the Fortune 100 use Gemini Enterprise, and 18 Cloud product lines each generate more "
               "than $1B of annual revenue. On the call, analysts (Mark Shmulik, Ross Sandler, Justin Post) asked how "
               "much of the growth is TPU hardware, which likely carries lower margin and less recurring revenue. "
               "Management did not quantify it. Without that split, we treat part of the 2027 Cloud growth as "
               "lower-quality hardware revenue."])
    figure(doc, "Google Cloud revenue ($B) and operating margin (%), Q4'24-Q3'26", figs["cloud"],
           [("link", "Alphabet 8-K Ex. 99.1 releases", URL["q3_pr"]), "; our calculations."])
    figure(doc, "Google Cloud backlog ($B), Q3'25-Q3'26", figs["backlog"],
           [("link", "Q3'26 earnings call (Oct 28, 2026)", URL["call"]), " for $561B; Q4'25 and Q3'25 per ",
            ("link", "MarketBeat (Feb 2026)", URL["q4_backlog"]), "; Q1'26 per ",
            ("link", "S&P Global (May 2026)", URL["q1_capex"]), "; Q2'26 per ",
            ("link", "Fierce Network (Jul 2026)", URL["q2_capex"]),
            ". Prior-quarter figures come from secondary coverage of earlier calls."])

    subheading(doc, "AI usage, Search and YouTube")
    rich(doc, ["CEO Sundar Pichai said Alphabet is \"now at $0.5 trillion annual run rate\" for revenue. The release "
               "describes it as \"a half-trillion dollar revenue run rate.\" First-party model APIs process \"nearly "
               "32 billion tokens per minute,\" up 40% QoQ and more than 4x YoY. The Gemini app has 1 billion users, "
               "and Gemini 4 models (Argon and Carbon) were announced. In YouTube, Chief Business Officer Philipp "
               "Schindler said living-room viewing is now over one-third of YouTube ad revenue. Management cited "
               "about 150M YouTube Music and Premium subscriptions (Pichai's figure included trials; Schindler's "
               "described Premium alone, so we do not treat these as a clean metric). On the call, Doug Anmuth "
               "(JPMorgan) asked about the gap between strong Search query growth and slowing Search revenue. We "
               "think it is the most important question for Services into 2027. Faster query growth and slower "
               "revenue growth suggest AI Mode is monetizing below traditional search for now, though FX and comps "
               "explain part of the gap. Paid-click and CPC data were not disclosed."])

    subheading(doc, "Guidance: 2026 capex held at $195-205B; 2027 to rise significantly")
    rich(doc, ["Alphabet gives no revenue or EPS guidance. Capex guidance is the main forward marker. It started at "
               "$175-185B in February, rose to $180-190B in April and to $195-205B in July, and was held in October. "
               "About 65% of Q3 technical infrastructure spend went to servers and networking (generally six-year "
               "useful lives), with 35% to data centers. Management repeated that 2027 capex will \"increase "
               "significantly\" and will detail it on the Q4 call. Future purchase commitments reached $924B, although "
               "management noted that many have variable pricing and volume flexibility. Our take: holding 2026 was "
               "expected. The bigger issue is that 2027 spending is open-ended while TTM FCF has fallen to $19.7B. "
               "We now model 2027 capex at $270B (+35%), up from $250B. That is above our prior view but, in our "
               "judgment, consistent with \"significantly.\""])
    figure(doc, "Evolution of 2026 capex guidance ($B)", figs["guide"],
           [("link", "Q4'25 release (Feb 4, 2026)", URL["q4_pr"]), "; April raise per ",
            ("link", "S&P Global", URL["q1_capex"]), "; July raise per ",
            ("link", "Fierce Network", URL["q2_capex"]), "; October per ",
            ("link", "Q3'26 call transcript", URL["call"]), "."])
    figure(doc, "Operating cash flow, capex and free cash flow ($B), Q4'24-Q3'26", figs["cash"],
           [("link", "Alphabet 8-K Ex. 99.1 releases", URL["q3_pr"]),
            "; FCF = OCF less purchases of property and equipment (derived where not reported)."])

    # ---------------- THESIS ----------------
    heading(doc, "Updated Investment Thesis")
    bullet(doc, "Pillar 1 - Google Cloud becomes a second profit engine. Status: STRENGTHENED.",
           ["We expected Cloud to be the main source of operating income growth from 2026. Q3 confirmed it. Cloud "
            "operating income tripled to $11.1B and supplied 80% of the YoY increase in total operating income "
            "($7.5B of $9.3B). Backlog of $561B covers more than four years of current Cloud revenue. The one "
            "caveat is quality of mix. TPU system sales are lumpier and probably lower margin than GCP services, "
            "and management expects third-party compute to pressure margins near term. We raise FY27E Cloud revenue "
            "6% to $191.6B but cut the margin assumption to 31% from 33%."])
    bullet(doc, "Pillar 2 - Search holds up through the AI transition. Status: UNCHANGED (watch closely).",
           ["Search & other grew 12%, the slowest of the last four quarters. The bear case says AI answers erode monetization, "
            "and Q3 neither proves nor disproves it. In support: query growth was the best in five years, AI Mode "
            "queries are 3x longer, and management attributes most of the slowdown to FX and comps. Against: "
            "revenue growth fell five points QoQ, the Q4 comp is harder, and no click/CPC data were given. We cut "
            "FY27E Services revenue 2% and will look for Search growth to stabilize around 10% in 1H'27."])
    bullet(doc, "Pillar 3 - Full-stack AI (TPUs + Gemini) is a cost and product moat. Status: STRENGTHENED.",
           ["Token volumes up 4x YoY, 1B Gemini app users and external TPU sales show the vertical stack is working "
            "both as a product and as a business. Selling TPU systems to third parties, including competitors (a "
            "topic Michael Nathanson raised on the call), adds a merchant-silicon revenue stream few peers have."])
    bullet(doc, "Pillar 4 - Capital returns are funded from free cash flow. Status: WEAKENED.",
           ["This is the pillar Q3 damaged. FCF was negative for a second quarter (-$5.9B in Q2, -$9.1B in Q3). "
            "Buybacks stopped in 2026 (vs. $39.8B in 9M'25). Alphabet raised $31.1B of notes in Q1, $30.5B of "
            "common equity and $19.1B of mandatory convertible preferred in Q2, and $12.4B through the ATM in Q3, "
            "with $27.6B left on the program. Long-term debt is $125.1B (vs. $46.5B at year-end). Net cash is still "
            "about $140B, so the balance sheet is not stretched, but holders are being diluted to fund capex. Our "
            "BUY now rests on the return on that capex, not on capital returns."])

    subheading(doc, "Risks update")
    rich(doc, [("b", "Worsened: "), "(1) Capex/FCF. Each additional $25B of 2027 capex lowers FY27E FCF by the same "
               "amount, and depreciation (+49% YoY in Q3) will rise faster than revenue into 2027. (2) Dilution and "
               "financing costs from ATM issuance, preferred dividends and interest expense, which we estimate cost "
               "~$0.09 of EPS per quarter. (3) Earnings volatility from $90.4B of marketable equity securities "
               "(SpaceX and others). ", ("b", "New: "), "Cloud margin pressure from third-party compute, and the "
               "lower margin and lumpiness of TPU system sales. ", ("b", "Unchanged: "),
               "Regulatory and antitrust remedies (the $5.2B EC fine was paid in Q3), Search disruption from AI "
               "assistants, and Other Bets losses ($2.0B in Q3). ", ("b", "Mitigated: "),
               "Cloud demand uncertainty. A $561B backlog with 55%+ converting in 24 months reduces it."])

    # ---------------- VALUATION ----------------
    heading(doc, "Valuation & Estimates")
    subheading(doc, "Price target methodology: $360 (50% P/E, 50% SOTP)")
    rich(doc, ["We set a 12-month price target of $360. It is the average of two methods ($357), rounded to the nearest $10. ",
               ("b", "(1) P/E: "), "28x our 2027E EPS ex-marks of $12.72 = $356. 28x is a premium to the 22.9x the "
               "stock trades at pre-market. We think a premium is justified by 24% revenue growth in 2027E and a "
               "Cloud segment growing 75%. It also sits below the ~34x implied by the $430 median Street target on "
               "our EPS. ", ("b", "(2) Sum-of-the-parts (2027E): "), "Google Services NOPAT of $141B at 22x ($3.11T). "
               "Google Cloud NOPAT of $48B at 35x ($1.69T), reflecting 75% growth and backlog visibility. Other Bets "
               "at a nominal $50B. Alphabet-level costs capitalized at 22x (-$0.54T). Net cash of $140B ($265.1B "
               "cash and marketable securities less $125.1B long-term debt). That totals $4.46T, or $358 per share "
               "on 12,450M diluted shares (which includes expected ATM and preferred dilution). NOPAT uses our 18.5% "
               "tax-rate estimate. ", ("b", "Sensitivity: "), "every 1x of P/E is worth about $13 per share. At 24x "
               "the P/E method gives $305; at 32x, $407. ", ("b", "Scenarios: "),
               "Bull $430 (32x on $13.50 if TPU sales and Search re-accelerate). Bear $265 (22x on $12.00 if 2027 "
               "capex exceeds $300B and Cloud margins fall below 28%)."])
    figure(doc, "Sum-of-the-parts valuation ($ per share, 2027E)", figs["sotp"],
           ["Our estimates; balance sheet from the ", ("link", "Q3'26 release", URL["q3_pr"]),
            " and ", ("link", "Q3'26 call", URL["call"]), "; pre-market price per ",
            ("link", "StockAnalysis (Oct 29, 2026)", URL["sa_quote"]), "."])

    subheading(doc, "What changed in our estimates")
    rich(doc, ["FY26E revenue edges down 0.2% to $495.7B. Q3 Cloud upside (+$1.7B vs. our model) and a higher Q4 Cloud "
               "forecast ($34.0B) offset a weaker Services outlook (Q4 Services $104.0B, +8.5%). FY26E operating "
               "income falls 1.9% to $163.9B because we lower Q4 Cloud margin to 32%. FY26E EPS ex-marks falls 5.2% "
               "to $10.98. That reflects the Q3 shortfall (-$0.33 vs. our model) and higher interest and preferred "
               "costs in Q4. Our Q4'26E EPS is $2.83. For FY27E, revenue rises 0.6% to $616.2B (Cloud +75% to "
               "$191.6B; Services +10% to $422.9B). Operating income falls 2.6% to $194.3B (31.5% margin) on heavier "
               "depreciation and third-party compute costs. EPS ex-marks falls 3.0% to $12.72 after we add $1.2B of "
               "preferred dividends and raise diluted shares to 12,450M. We raise 2027E capex to $270B and now "
               "forecast FY27E FCF of about -$40B. Our EPS is well below the $15.19 FY27 Street figure, which may "
               "include different treatment of investment gains. We would rather start from a conservative base."])
    figure(doc, "Estimate revisions: prior vs. new, FY26E and FY27E", figs["rev"],
           ["Our estimates (prior = pre-print model); FY26E incorporates 9M'26 actuals from the ",
            ("link", "Q3'26 release", URL["q3_pr"]), "."])

    subheading(doc, "Bottom line")
    rich(doc, ["The quarter answered the question of whether AI infrastructure spending earns a return. In Cloud it "
               "clearly does: revenue doubled and margin reached 36%. It also raised a new question: how long "
               "shareholders will fund a capex cycle that has turned FCF negative and replaced buybacks with "
               "issuance. With the stock down ~10% on a headline EPS miss that was mostly a non-cash SpaceX mark, we "
               "think the risk/reward favors buyers. We maintain BUY with a $360 PT. Catalysts: Q4 results and 2027 "
               "capex guidance (early February 2027), Cloud margin through the third-party compute ramp, and Search "
               "growth once FX and World Cup comps ease."])

    # ---------------- SOURCES ----------------
    heading(doc, "Sources & References")
    src = [
        ("Earnings materials (Q3 2026)", [
            ("Alphabet Q3 2026 earnings release, Form 8-K Exhibit 99.1 (filed October 28, 2026)", URL["q3_pr"]),
            ("Form 8-K filing index, accession 0001652044-26-000080 (October 28, 2026)", URL["q3_8k"]),
            ("Form 10-Q for quarter ended September 30, 2026 (accession 0001652044-26-000084, filed October 29, 2026)", URL["q3_10q"]),
            ("Q3 2026 earnings call transcript, October 28, 2026 (The Motley Fool)", URL["call"]),
            ("Q3 2026 earnings call transcript and summary, October 28, 2026 (Investing.com)", URL["call_inv"]),
            ("Q3 2026 earnings slides summary (Investing.com, October 28, 2026)", URL["slides_inv"]),
        ]),
        ("Prior-period materials", [
            ("Q2 2026 earnings release, 8-K Exhibit 99.1 (July 22, 2026)", URL["q2_pr"]),
            ("Q1 2026 earnings release, 8-K Exhibit 99.1 (April 29, 2026)", URL["q1_pr"]),
            ("Q4 2025 earnings release, 8-K Exhibit 99.1 (February 4, 2026) - initial 2026 capex guidance", URL["q4_pr"]),
            ("S&P Global Market Intelligence, Alphabet post-Q1'26 snapshot (May 15, 2026) - Q1 capex raise, $462B backlog", URL["q1_capex"]),
            ("Fierce Network, \"Google Cloud surges, but Alphabet's AI capex problem grows\" (July 2026) - Q2 capex raise, $514B backlog", URL["q2_capex"]),
            ("MarketBeat, Q4 2025 results coverage (February 2026) - $240B Q4'25 and $155B Q3'25 backlog", URL["q4_backlog"]),
            ("SEC EDGAR: Alphabet 10-Q filing list", URL["edgar"]),
        ]),
        ("Consensus, pricing and market data", [
            ("MarketBeat, Alphabet Q3 2026 earnings report page - consensus EPS $3.01, revenue $127.97B", URL["mb"]),
            ("CNBC, Alphabet Q3 2026 earnings coverage (October 28, 2026) - $3.02 EPS consensus cited", URL["cnbc"]),
            ("StockAnalysis, GOOGL price history - October 28, 2026 close $324.89", URL["sa_hist"]),
            ("StockAnalysis, GOOGL quote - pre-market $291.08 (6:15 AM ET, October 29, 2026), 52-week range", URL["sa_quote"]),
            ("StockAnalysis, GOOGL analyst forecast - consensus PT $429.68, FY26/FY27 EPS, analyst actions", URL["sa_fc"]),
        ]),
    ]
    for group, items in src:
        para(doc, group, size=10, bold=True, space_before=3, space_after=2, keep_next=True)
        for text, url in items:
            p = para(doc, size=9.5, space_after=1)
            set_run_font(p.add_run("• "), 9.5)
            add_hyperlink(p, text, url, size=9.5)
    para(doc, "", space_after=2)
    rich(doc, [("i", "Disclosures: Estimates, price target and ratings are the analyst's own and are not company "
                     "guidance. \"EPS ex-marks\" is a non-GAAP analyst measure that removes gains and losses on equity "
                     "securities at an assumed 21% tax rate. Prior-quarter Cloud backlog and capex-guidance history come "
                     "from secondary press coverage of earlier earnings calls and were not checked against those "
                     "transcripts. This report is for informational purposes only and is not investment advice.")],
         size=8.5)

    doc.save(OUT)
    print("saved", OUT, "figures:", FIG[0], "tables:", TBL[0])


if __name__ == "__main__":
    build()
